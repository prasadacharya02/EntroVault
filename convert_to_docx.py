# convert_to_docx.py
"""
Converts docs/ENTROVAULT_IEEE_PROJECT_REPORT_100_PAGES.md into a beautifully formatted,
professional academic Word Document (.docx) conforming to IEEE Dissertation standards.
Converts both pipe markdown tables and ASCII box tables into native Word tables.
"""

import os
import re
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def sanitize_text(text):
    """Remove control characters invalid in XML."""
    if not text:
        return ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding/margins for table cells in dxa (1/20th of a point)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_inline_formatted_text(paragraph, text):
    """Parse basic inline markdown (bold **text**, code `text`, italic *text*)."""
    text = sanitize_text(text)
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        token = sanitize_text(token)
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(144, 26, 30)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)

def create_native_table(doc, rows_data, title=None):
    """Construct a styled academic Word table."""
    if not rows_data:
        return

    if title:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(12)
        tp.paragraph_format.space_after = Pt(4)
        tp.paragraph_format.keep_with_next = True
        trun = tp.add_run(title)
        trun.bold = True
        trun.font.name = 'Calibri'
        trun.font.size = Pt(10.5)
        trun.font.color.rgb = RGBColor(11, 37, 69)

    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for i, row_data in enumerate(rows_data):
        row = table.rows[i]
        is_header = (i == 0)
        for j in range(num_cols):
            cell = row.cells[j]
            cell_text = row_data[j] if j < len(row_data) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            add_inline_formatted_text(p, cell_text)

            if is_header:
                set_cell_background(cell, "0B2545") # Navy
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                for r in p.runs:
                    r.bold = True
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg_color = "F7F9FB" if (i % 2 == 0) else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(30, 30, 30)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def md_to_docx():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(base_dir, "docs", "ENTROVAULT_IEEE_PROJECT_REPORT_100_PAGES.md")
    docx_path = os.path.join(base_dir, "docs", "ENTROVAULT_IEEE_PROJECT_REPORT_100_PAGES.docx")

    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found.")
        sys.exit(1)

    print(f"Reading markdown source: {md_path}")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Configure A4 page layout with 1-inch margins
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("EntroVault: Proactive Ransomware Defence | IEEE Major Project Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(120, 120, 120)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(20, 20, 20)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal in_code_block, code_block_lines
        if code_block_lines:
            raw_text = "".join(code_block_lines)
            stripped_lines = [l.strip() for l in raw_text.strip().split('\n')]

            # Check if this code block is actually an ASCII box table
            is_box_table = any("TABLE " in l for l in stripped_lines) and any(l.startswith("+") for l in stripped_lines)
            if is_box_table:
                # Parse box table
                title = None
                rows_data = []
                for sl in stripped_lines:
                    if sl.startswith("|") and not sl.startswith("+"):
                        cells = [c.strip() for c in sl.strip("|").split("|")]
                        if len(cells) == 1 and "TABLE " in cells[0]:
                            title = cells[0]
                        else:
                            rows_data.append(cells)
                if rows_data:
                    create_native_table(doc, rows_data, title=title)
                    code_block_lines = []
                    in_code_block = False
                    return

            # Otherwise, render as standard monospace code block
            code_text = sanitize_text(raw_text.rstrip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(code_text)
            run.font.name = 'Courier New'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
            code_block_lines = []
        in_code_block = False

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return

        rows = []
        for line in table_lines:
            parts = [c.strip() for c in line.strip().strip('|').split('|')]
            if parts and all(re.match(r'^:?-+:?$', p) for p in parts if p):
                continue
            rows.append(parts)

        if rows:
            create_native_table(doc, rows)

        table_lines = []
        in_table = False

    print("Parsing Markdown and populating Word Document structures...")
    idx = 0
    total_lines = len(lines)

    while idx < total_lines:
        line = lines[idx]

        # Check for page-break tag
        if '<div style="page-break-after: always;"></div>' in line:
            if in_code_block:
                flush_code_block()
            if in_table:
                flush_table()
            doc.add_page_break()
            idx += 1
            continue

        # Code fences
        if line.strip().startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_block_lines = []
            idx += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            idx += 1
            continue

        # Table rows
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            in_table = True
            table_lines.append(line)
            idx += 1
            continue
        elif in_table:
            flush_table()

        stripped = line.strip()

        if not stripped:
            idx += 1
            continue

        # Headings
        if stripped.startswith('# '):
            heading_text = sanitize_text(stripped[2:].strip())
            if heading_text.startswith("CHAPTER ") and len(doc.paragraphs) > 5:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(heading_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(17)
            run.bold = True
            run.font.color.rgb = RGBColor(11, 37, 69)
            idx += 1
            continue

        elif stripped.startswith('## '):
            heading_text = sanitize_text(stripped[3:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(heading_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(19, 60, 85)
            idx += 1
            continue

        elif stripped.startswith('### '):
            heading_text = sanitize_text(stripped[4:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(heading_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = RGBColor(29, 78, 137)
            idx += 1
            continue

        elif stripped.startswith('#### '):
            heading_text = sanitize_text(stripped[5:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(heading_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = RGBColor(60, 60, 60)
            idx += 1
            continue

        # Horizontal Rule
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 55)
            run.font.color.rgb = RGBColor(180, 180, 180)
            run.font.size = Pt(8)
            idx += 1
            continue

        # Blockquote
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sanitize_text(stripped[2:].strip()))
            run.italic = True
            run.font.color.rgb = RGBColor(60, 60, 60)
            idx += 1
            continue

        # Unordered lists
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, stripped[2:].strip())
            idx += 1
            continue

        # Ordered lists
        numbered_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if numbered_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, numbered_match.group(2))
            idx += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        add_inline_formatted_text(p, stripped)
        idx += 1

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    print(f"Saving compiled Word Document (.docx): {docx_path} ...")
    doc.save(docx_path)
    file_size_mb = os.path.getsize(docx_path) / (1024 * 1024)
    print(f"Word Document Successfully Generated!")
    print(f"Output File: {docx_path}")
    print(f"File Size:   {file_size_mb:.2f} MB ({os.path.getsize(docx_path):,} bytes)")

if __name__ == "__main__":
    md_to_docx()
